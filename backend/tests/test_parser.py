"""Unit tests for multi-format document parser (backed by Docling)."""

from unittest.mock import MagicMock, patch

import pytest

from ingestion.parser import parse_document, ParseResult


class TestParseTxt:
    def test_parse_plain_text(self, tmp_path):
        f = tmp_path / "sample.txt"
        f.write_text("Hello world.\nSecond line.", encoding="utf-8")
        result = parse_document(str(f))
        assert isinstance(result, ParseResult)
        assert "Hello world." in result.text
        assert "Second line." in result.text
        assert result.format == "txt"

    def test_parse_empty_text(self, tmp_path):
        f = tmp_path / "empty.txt"
        f.write_text("", encoding="utf-8")
        result = parse_document(str(f))
        assert result.text == ""
        assert result.format == "txt"


class TestParseMarkdown:
    def test_parse_markdown(self, tmp_path):
        f = tmp_path / "readme.md"
        content = (
            "# Title\n\n"
            "Some **bold** text.\n\n"
            "## Subtitle ##\n\n"
            "## C# ##\n\n"
            "```python\n"
            "# This is a code comment inside python block\n"
            "print('hello')\n"
            "```\n"
        )
        f.write_text(content, encoding="utf-8")
        result = parse_document(str(f))
        assert "Title" in result.text
        assert "bold" in result.text
        assert result.format == "md"
        assert result.headings == ["Title", "Subtitle", "C#"]


class TestParsePdf:
    def test_parse_pdf(self, tmp_path):
        """Create a minimal PDF with PyMuPDF and parse it back via Docling."""
        pytest.importorskip("fitz", reason="PyMuPDF not installed (test fixture only)")
        import fitz

        pdf_path = str(tmp_path / "test.pdf")
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "PDF content here")
        doc.save(pdf_path)
        doc.close()

        result = parse_document(pdf_path)
        assert "PDF content here" in result.text
        assert result.format == "pdf"
        assert result.page_count >= 1


class TestParseDocx:
    def test_parse_docx(self, tmp_path):
        """Create a minimal DOCX with python-docx and parse it back via Docling."""
        pytest.importorskip("docx", reason="python-docx not installed (test fixture only)")
        from docx import Document

        docx_path = str(tmp_path / "test.docx")
        doc = Document()
        doc.add_heading("Test Heading", level=1)
        doc.add_paragraph("Paragraph content.")
        doc.save(docx_path)

        result = parse_document(docx_path)
        assert "Test Heading" in result.text
        assert "Paragraph content." in result.text
        assert result.format == "docx"

    def test_parse_docx_with_table(self, tmp_path):
        pytest.importorskip("docx", reason="python-docx not installed (test fixture only)")
        from docx import Document

        docx_path = str(tmp_path / "table.docx")
        doc = Document()
        doc.add_heading("Doc with Table", level=1)

        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header 1"
        table.cell(0, 1).text = "Header 2"
        table.cell(1, 0).text = ""
        table.cell(1, 1).text = "Cell B"
        doc.save(docx_path)

        result = parse_document(docx_path)
        assert "Doc with Table" in result.text
        # Table content should appear in some form
        assert "Header 1" in result.text
        assert "Header 2" in result.text
        assert "Cell B" in result.text


class TestParseHtml:
    def test_parse_html(self, tmp_path):
        f = tmp_path / "page.html"
        f.write_text(
            "<html><body><h1>Header</h1><p>Content here</p></body></html>",
            encoding="utf-8",
        )
        result = parse_document(str(f))
        assert "Header" in result.text
        assert "Content here" in result.text
        assert result.format == "html"

    def test_parse_htm(self, tmp_path):
        f = tmp_path / "page.htm"
        f.write_text(
            "<html><body><h1>Header HTM</h1><p>Content HTM</p></body></html>",
            encoding="utf-8",
        )
        result = parse_document(str(f))
        assert "Header HTM" in result.text
        assert "Content HTM" in result.text
        # htm is normalised to html
        assert result.format == "html"

    def test_strips_scripts_and_styles(self, tmp_path):
        f = tmp_path / "messy.html"
        f.write_text(
            "<html><head><style>body{}</style></head>"
            "<body><script>alert(1)</script><p>Clean text</p></body></html>",
            encoding="utf-8",
        )
        result = parse_document(str(f))
        assert "Clean text" in result.text
        assert "alert" not in result.text
        assert "body{}" not in result.text


class TestParseImage:
    def test_parse_image_via_docling(self, tmp_path):
        """Verify that image files are routed through Docling's OCR pipeline."""
        pytest.importorskip("PIL", reason="Pillow not installed (test fixture only)")
        from PIL import Image

        img_path = tmp_path / "dummy.png"
        img = Image.new("RGB", (100, 100), color="white")
        img.save(img_path)

        # Mock the DocumentConverter so we don't need a real OCR engine in CI
        mock_doc = MagicMock()
        mock_doc.export_to_markdown.return_value = "Mocked OCR Text from Image"
        mock_doc.pages = [MagicMock()]
        mock_doc.iterate_items.return_value = []

        mock_result = MagicMock()
        mock_result.document = mock_doc

        with patch(
            "docling.document_converter.DocumentConverter.convert",
            return_value=mock_result,
        ):
            result = parse_document(str(img_path))

        assert "Mocked OCR Text from Image" in result.text
        assert result.format == "png"


class TestCleanText:
    def test_deduplicates_whitespace(self, tmp_path):
        f = tmp_path / "spaces.txt"
        f.write_text("Too   many    spaces\n\n\n\nand lines.", encoding="utf-8")
        result = parse_document(str(f))
        assert result.text == "Too many spaces\n\nand lines."

    def test_strips_leading_trailing(self, tmp_path):
        f = tmp_path / "padded.txt"
        f.write_text("   padded content   ", encoding="utf-8")
        result = parse_document(str(f))
        assert result.text == "padded content"

    def test_preserves_indentation(self, tmp_path):
        f = tmp_path / "indented.txt"
        f.write_text("Header\n  list item 1\n    nested list item\n", encoding="utf-8")
        result = parse_document(str(f))
        assert result.text == "Header\n  list item 1\n    nested list item"


class TestFileNotFound:
    def test_file_not_found_raises_error(self):
        with pytest.raises(FileNotFoundError, match="File not found"):
            parse_document("nonexistent_file_path.txt")

    def test_directory_path_raises_error(self, tmp_path):
        with pytest.raises(ValueError, match="Path is not a file"):
            parse_document(str(tmp_path))


class TestUnsupportedFormat:
    def test_unsupported_raises_error(self, tmp_path):
        f = tmp_path / "data.xyz"
        f.write_text("binary stuff", encoding="utf-8")
        with pytest.raises(ValueError, match="Unsupported"):
            parse_document(str(f))
